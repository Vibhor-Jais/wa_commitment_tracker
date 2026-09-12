import io

from openpyxl import load_workbook

from wa_commitment_tracker.builder import build_dataframe, build_workbook
from wa_commitment_tracker.parser import FIELDS

EXPORT = """12/09/2026, 08:00 - Alice: Good morning Team,
CA noa:- 
SA noa:- 1
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 5L
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
12/09/2026, 18:30 - Alice: Achievement update,
CA noa:- 
SA noa:- 
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 5L
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
12/09/2026, 09:15 - Bob: Good morning,
CA noa:- 
SA noa:- 1
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
"""


def _build():
    df, warnings = build_dataframe(EXPORT)
    buf = io.BytesIO()
    build_workbook(df, buf)
    buf.seek(0)
    return load_workbook(buf)


def test_header_layout_has_two_columns_per_field():
    wb = _build()
    ws = wb["12-09-2026"]
    assert ws["A1"].value == "Member"
    sa_value_col_idx = 2 + FIELDS.index("SA value") * 2  # 1-indexed: B=2, each field takes 2 cols
    from openpyxl.utils import get_column_letter
    commit_col = get_column_letter(sa_value_col_idx)
    achieve_col = get_column_letter(sa_value_col_idx + 1)
    assert ws[f"{commit_col}1"].value == "SA value"
    assert ws[f"{commit_col}2"].value == "Commitment"
    assert ws[f"{achieve_col}2"].value == "Achievement"


def test_one_row_per_member_with_both_types_side_by_side():
    wb = _build()
    ws = wb["12-09-2026"]
    from openpyxl.utils import get_column_letter
    sa_value_idx = 2 + FIELDS.index("SA value") * 2
    commit_col = get_column_letter(sa_value_idx)
    achieve_col = get_column_letter(sa_value_idx + 1)

    # exactly 2 data rows (Alice, Bob) - Alice's commitment+achievement merged into ONE row
    member_names = [ws.cell(row=r, column=1).value for r in range(3, 5)]
    assert member_names == ["Alice", "Bob"]

    alice_row = 3
    bob_row = 4
    assert ws[f"{commit_col}{alice_row}"].value == 500_000
    assert ws[f"{achieve_col}{alice_row}"].value == 500_000
    assert ws[f"{commit_col}{bob_row}"].value is None  # Bob never mentioned SA value
    assert ws[f"{achieve_col}{bob_row}"].value is None


def test_total_row_and_daily_summary_present():
    wb = _build()
    ws = wb["12-09-2026"]
    labels = [ws.cell(row=r, column=1).value for r in range(1, 10)]
    assert "Total" in labels
    summary = wb["Daily Summary"]
    assert summary["A1"].value == "Date"
    assert summary["B1"].value == "Present (Commitment)"
    assert summary["C1"].value == "Present (Achievement)"
