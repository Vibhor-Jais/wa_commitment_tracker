import io

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from wa_commitment_tracker.builder import update_workbook
from wa_commitment_tracker.parser import FIELDS

DAY1 = """10/09/2026, 08:00 - Alice: Good morning,
CA noa:- 
SA noa:- 1
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 2L
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
"""

DAY1_AND_2 = DAY1 + """11/09/2026, 08:00 - Alice: Good morning,
CA noa:- 
SA noa:- 1
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 3L
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
"""

DAY1_2_AND_3 = DAY1_AND_2 + """12/09/2026, 08:00 - Alice: Good morning,
CA noa:- 
SA noa:- 1
SA Affluent -
Elite:- 
M0 Value:- 
Ca value:- 
Sa value:- 4L
RTD:- 
RD:- 
LI:- 
HI:- 
MF:-
SIP:- 
PMJDY:-
APY:- 
PMJJBY:- 
PMSBY:-
"""


def _save_to_buffer(export_text, existing_buf=None):
    out = io.BytesIO()
    added, skipped, warnings = update_workbook(export_text, existing_buf, out)
    out.seek(0)
    return added, skipped, warnings, out


def test_first_run_creates_all_dates_with_none_existing():
    added, skipped, warnings, buf = _save_to_buffer(DAY1_AND_2, existing_buf=None)
    assert warnings == []
    assert skipped == []
    assert set(added) == {"10/09/2026", "11/09/2026"}
    wb = load_workbook(buf)
    assert "10-09-2026" in wb.sheetnames
    assert "11-09-2026" in wb.sheetnames


def test_second_run_only_adds_new_date():
    _, _, _, buf1 = _save_to_buffer(DAY1_AND_2, existing_buf=None)

    # re-run with a cumulative export that now includes day 3, feeding
    # back in the workbook produced by the first run
    added, skipped, warnings, buf2 = _save_to_buffer(DAY1_2_AND_3, existing_buf=buf1)

    assert added == ["12/09/2026"]
    assert set(skipped) == {"10/09/2026", "11/09/2026"}

    wb = load_workbook(buf2)
    assert set(wb.sheetnames) >= {"10-09-2026", "11-09-2026", "12-09-2026", "Daily Summary"}
    # existing day-1 sheet must be untouched: still shows the original 2L value
    ws1 = wb["10-09-2026"]
    sa_value_commit_col = get_column_letter(2 + FIELDS.index("SA value") * 2)
    assert ws1[f"{sa_value_commit_col}3"].value == 200_000

    summary = wb["Daily Summary"]
    dates_in_summary = [summary.cell(row=r, column=1).value for r in range(2, summary.max_row + 1)]
    assert dates_in_summary == ["10/09/2026", "11/09/2026", "12/09/2026"]


def test_rerunning_with_nothing_new_adds_nothing():
    _, _, _, buf1 = _save_to_buffer(DAY1_AND_2, existing_buf=None)
    added, skipped, warnings, buf2 = _save_to_buffer(DAY1_AND_2, existing_buf=buf1)
    assert added == []
    assert set(skipped) == {"10/09/2026", "11/09/2026"}
    wb = load_workbook(buf2)
    # still exactly 2 date sheets, no duplicates created
    date_sheets = [n for n in wb.sheetnames if n not in ("Daily Summary", "Notes")]
    assert len(date_sheets) == 2
